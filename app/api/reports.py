"""财务报表：利润表、现金流量表、账户余额表、科目余额表、资产负债简表、经营分析、Excel 导出。"""
from __future__ import annotations

import io
from datetime import date, timedelta

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from .. import db, rules
from ..api.finance import MONEY_KINDS, _balance
from ..deps import require

router = APIRouter(prefix="/api/rep", tags=["reports"])

INCOME_GROUPS = ["主营业务", "其他收益", "资金往来", "其他"]
EXPENSE_GROUPS = ["饲料", "防疫医疗", "繁殖", "生产费用", "管理费用", "资金往来", "其他"]

# 现金流量表口径：只统计资金账户的实收实付，剔除内部对转/提取现金这类「钱换个口袋」的腿
MONEY_K = "a.kind IN ('cash','bank','wechat','alipay','other')"
CF_KEEP = "t.link_type <> 'transfer' AND c.name NOT IN ('内部转账','提取现金')"


class RangeIn(BaseModel):
    start: str
    end: str


def _range(start: str | None, end: str | None, preset: str = "") -> tuple[str, str]:
    if start and end:
        return start, end
    if preset:
        return rules.period_range(preset)
    return rules.period_range("month")


def _pl(start: str, end: str) -> dict:
    """损益汇总：按科目分组。affect=0 的往来科目不计入。"""
    inc = db.rows(
        "SELECT c.group_name AS g, c.name AS n, SUM(t.amount) AS v, COUNT(*) AS cnt"
        " FROM txn t JOIN categories c ON c.id=t.category_id"
        " WHERE t.biz_date BETWEEN ? AND ? AND t.direction='income' AND t.affect=1"
        " GROUP BY c.group_name, c.name ORDER BY v DESC", (start, end))
    exp = db.rows(
        "SELECT c.group_name AS g, c.name AS n, SUM(t.amount) AS v, COUNT(*) AS cnt"
        " FROM txn t JOIN categories c ON c.id=t.category_id"
        " WHERE t.biz_date BETWEEN ? AND ? AND t.direction='expense' AND t.affect=1"
        " GROUP BY c.group_name, c.name ORDER BY v DESC", (start, end))

    def pack(rs):
        groups: dict[str, list] = {}
        for r in rs:
            groups.setdefault(r["g"] or "其他", []).append(
                {"name": r["n"], "amount": round(r["v"] or 0, 2), "count": r["cnt"]})
        out = [{"group": g, "items": items, "amount": round(sum(i["amount"] for i in items), 2)}
               for g, items in groups.items()]
        out.sort(key=lambda x: -x["amount"])
        return out

    total_inc = sum(g["amount"] for g in pack(inc))
    total_exp = sum(g["amount"] for g in pack(exp))
    return {
        "income": pack(inc), "expense": pack(exp),
        "total_income": round(total_inc, 2), "total_expense": round(total_exp, 2),
        "profit": round(total_inc - total_exp, 2),
        "margin": round((total_inc - total_exp) / total_inc * 100, 1) if total_inc else 0,
    }


@router.get("/profit")
def profit(start: str = "", end: str = "", preset: str = "", compare: bool = True,
           _: dict = Depends(require("rep.read"))):
    s, e = _range(start, end, preset)
    cur = _pl(s, e)
    out = {"start": s, "end": e, "current": cur, "compare": None}
    if compare:
        d = rules.parse_date(s)
        span = (rules.parse_date(e) - d).days + 1
        ps = (d - timedelta(days=span)).isoformat()
        pe = (d - timedelta(days=1)).isoformat()
        if ps >= "2000-01-01":
            prev = _pl(ps, pe)
            out["compare"] = {
                "start": ps, "end": pe,
                "income": prev["total_income"], "expense": prev["total_expense"],
                "profit": prev["profit"],
                "income_delta": round(cur["total_income"] - prev["total_income"], 2),
                "expense_delta": round(cur["total_expense"] - prev["total_expense"], 2),
                "profit_delta": round(cur["profit"] - prev["profit"], 2),
            }
    return out


@router.get("/trend")
def trend(start: str = "", end: str = "", _: dict = Depends(require("rep.read"))):
    """逐月损益趋势。"""
    s, e = _range(start, end, "year")
    rs = db.rows(
        "SELECT strftime('%Y-%m', biz_date) AS m,"
        " SUM(CASE WHEN direction='income' THEN amount ELSE 0 END) AS inc,"
        " SUM(CASE WHEN direction='expense' THEN amount ELSE 0 END) AS exp"
        " FROM txn WHERE biz_date BETWEEN ? AND ? AND affect=1 GROUP BY m ORDER BY m", (s, e))
    out = [{"month": r["m"], "income": round(r["inc"] or 0, 2),
            "expense": round(r["exp"] or 0, 2),
            "profit": round((r["inc"] or 0) - (r["exp"] or 0), 2)} for r in rs]
    total = {"income": sum(o["income"] for o in out), "expense": sum(o["expense"] for o in out)}
    total["profit"] = round(total["income"] - total["expense"], 2)
    return {"start": s, "end": e, "items": out, "total": total}


@router.get("/cashflow")
def cashflow(start: str = "", end: str = "", preset: str = "", _: dict = Depends(require("rep.read"))):
    """现金流量表：按月 + 按账户。
    口径 = 资金账户实收实付（含借款还款、应收应付收付），
    剔除内部对转 / 提取现金这类「钱换个口袋」的动作。"""
    s, e = _range(start, end, preset)
    months = rules.month_list(s, e)
    by_month = db.rows(
        "SELECT strftime('%Y-%m', biz_date) AS m,"
        " SUM(CASE WHEN direction='income' THEN amount ELSE 0 END) AS inc,"
        " SUM(CASE WHEN direction='expense' THEN amount ELSE 0 END) AS exp"
        " FROM txn t JOIN accounts a ON a.id=t.account_id JOIN categories c ON c.id=t.category_id"
        f" WHERE t.biz_date BETWEEN ? AND ? AND {MONEY_K} AND {CF_KEEP}"
        " GROUP BY m ORDER BY m", (s, e))
    m_map = {r["m"]: r for r in by_month}
    rows, c_in, c_out = [], 0.0, 0.0
    for m in months:
        r = m_map.get(m)
        i = round(r["inc"] or 0, 2) if r else 0
        o = round(r["exp"] or 0, 2) if r else 0
        c_in += i; c_out += o
        rows.append({"month": m, "inflow": i, "outflow": o, "net": round(i - o, 2)})

    by_cat = db.rows(
        "SELECT c.name AS cat, t.direction AS dir, SUM(t.amount) AS v"
        " FROM txn t JOIN categories c ON c.id=t.category_id JOIN accounts a ON a.id=t.account_id"
        f" WHERE t.biz_date BETWEEN ? AND ? AND {MONEY_K} AND {CF_KEEP}"
        " GROUP BY c.name, t.direction ORDER BY v DESC", (s, e))
    by_acct = db.rows(
        "SELECT a.name AS acct, a.kind AS kind,"
        " SUM(CASE WHEN t.direction='income' THEN t.amount ELSE 0 END) AS inc,"
        " SUM(CASE WHEN t.direction='expense' THEN t.amount ELSE 0 END) AS exp"
        " FROM txn t JOIN accounts a ON a.id=t.account_id JOIN categories c ON c.id=t.category_id"
        f" WHERE t.biz_date BETWEEN ? AND ? AND {MONEY_K} AND {CF_KEEP}"
        " GROUP BY a.name, a.kind ORDER BY (inc-exp) DESC", (s, e))

    return {
        "start": s, "end": e, "rows": rows,
        "total_inflow": round(c_in, 2), "total_outflow": round(c_out, 2),
        "net": round(c_in - c_out, 2),
        "by_category": [{"category": r["cat"], "direction": r["dir"],
                         "amount": round(r["v"] or 0, 2)} for r in by_cat],
        "by_account": [{"account": r["acct"], "inflow": round(r["inc"] or 0, 2),
                        "outflow": round(r["exp"] or 0, 2),
                        "net": round((r["inc"] or 0) - (r["exp"] or 0), 2)} for r in by_acct],
    }


@router.get("/account-balance")
def account_balance(as_of: str = "", _: dict = Depends(require("rep.read"))):
    d = as_of or rules.today_str()
    items = [_balance(r["id"], d) for r in db.rows("SELECT * FROM accounts WHERE active=1")]
    money = [i for i in items if i["is_money"]]
    return {
        "date": d, "items": items,
        "money_total": round(sum(i["balance"] for i in money), 2),
        "receivable": round(sum(i["balance"] for i in items if i["kind"] == "receivable"), 2),
        "payable": round(sum(i["balance"] for i in items if i["kind"] == "payable"), 2),
        "grand_total": round(sum(i["balance"] for i in items), 2),
    }


@router.get("/category-ledger")
def category_ledger(start: str = "", end: str = "", preset: str = "",
                    _: dict = Depends(require("rep.read"))):
    s, e = _range(start, end, preset)
    rs = db.rows(
        "SELECT c.id, c.name, c.kind, c.group_name,"
        " SUM(CASE WHEN t.direction='income' THEN t.amount ELSE 0 END) AS inc,"
        " SUM(CASE WHEN t.direction='expense' THEN t.amount ELSE 0 END) AS exp,"
        " COUNT(*) AS cnt FROM categories c LEFT JOIN txn t"
        " ON t.category_id=c.id AND t.biz_date BETWEEN ? AND ? AND t.affect=1"
        " WHERE c.affect=1 GROUP BY c.id ORDER BY c.kind DESC, c.sort", (s, e))
    items, ti, te = [], 0.0, 0.0
    for r in rs:
        if not r["cnt"]:
            continue
        inc, exp = round(r["inc"] or 0, 2), round(r["exp"] or 0, 2)
        if r["kind"] == "income":
            ti += inc
        else:
            te += exp
        items.append({"id": r["id"], "name": r["name"], "kind": r["kind"],
                      "group_name": r["group_name"], "income": inc, "expense": exp,
                      "count": r["cnt"]})
    return {"start": s, "end": e, "items": items,
            "total_income": round(ti, 2), "total_expense": round(te, 2),
            "total_profit": round(ti - te, 2)}


@router.get("/balance-sheet")
def balance_sheet(as_of: str = "", _: dict = Depends(require("rep.read"))):
    """资产负债简表（管理用，非法定报表）。

    只呈现三件事：有什么资产、欠什么钱、还剩多少净资产，外加累计经营利润作参考。
    不做「资产=负债+权益」的强行配平——只有从建账第一天起每笔业务都录进系统，
    这个等式才会自然成立，中途开始用就会出现缺口，那样报出来反而误导。
    """
    d = as_of or rules.today_str()
    accs = db.rows("SELECT * FROM accounts WHERE active=1")
    cash = sum(_balance(a["id"], d)["balance"] for a in accs if a["kind"] in MONEY_KINDS)
    recv = sum(_balance(a["id"], d)["balance"] for a in accs if a["kind"] == "receivable")
    pay = sum(_balance(a["id"], d)["balance"] for a in accs if a["kind"] == "payable")

    # 未偿还借款 = 借过多少 - 还了多少
    loan = db.scalar(
        "SELECT COALESCE(SUM(CASE WHEN c.name='借款' THEN t.amount ELSE 0 END)"
        " - SUM(CASE WHEN c.name='还款' THEN t.amount ELSE 0 END),0)"
        " FROM txn t JOIN categories c ON c.id=t.category_id"
        " WHERE t.biz_date<=? AND t.affect=0 AND t.direction IN ('income','expense')", (d,)) or 0

    livestock = db.scalar(
        "SELECT COALESCE(SUM(COALESCE(buy_price,0)),0) FROM goat WHERE status='alive'", (), 0) or 0
    heads = db.scalar("SELECT COUNT(*) FROM goat WHERE status='alive'", (), 0)

    year_pl = _pl(f"{d[:4]}-01-01", d)
    all_pl = _pl("2000-01-01", d)

    assets = cash + recv + livestock
    liabilities = pay + loan
    return {
        "date": d,
        "assets": [
            {"name": "货币资金", "amount": round(cash, 2), "note": "现金 + 银行 + 微信支付宝"},
            {"name": "应收账款", "amount": round(recv, 2), "note": "客户欠款，尚未收回"},
            {"name": "生产性生物资产（存栏）", "amount": round(livestock, 2),
             "note": f"按购入成本计，在场 {heads} 只（自繁羔羊按 0 计）"},
        ],
        "liabilities": [
            {"name": "应付账款", "amount": round(pay, 2), "note": "供应商欠款，尚未支付"},
            {"name": "未偿还借款", "amount": round(loan, 2), "note": "借款累计减已还款累计"},
        ],
        "totals": {
            "assets": round(assets, 2),
            "liabilities": round(liabilities, 2),
            "net_asset": round(assets - liabilities, 2),
            "year_profit": round(year_pl["profit"], 2),
            "cum_profit": round(all_pl["profit"], 2),
            "heads": heads,
        },
    }


@router.get("/operations")
def operations(start: str = "", end: str = "", preset: str = "", _: dict = Depends(require("rep.read"))):
    """经营分析：奶业 + 存栏 + 单位效益。"""
    s, e = _range(start, end, preset)
    pl = _pl(s, e)
    milk_amt = db.scalar(
        "SELECT COALESCE(SUM(weight_kg*price),0) FROM milk WHERE biz_date BETWEEN ? AND ?", (s, e), 0) or 0
    milk_kg = db.scalar(
        "SELECT COALESCE(SUM(weight_kg),0) FROM milk WHERE biz_date BETWEEN ? AND ?", (s, e), 0) or 0
    days = db.scalar("SELECT COUNT(DISTINCT biz_date) FROM milk WHERE biz_date BETWEEN ? AND ?",
                     (s, e), 0) or 0
    milk_heads = db.scalar(
        "SELECT COUNT(*) FROM goat WHERE status='alive' AND sex='f' AND first_kidding IS NOT NULL", (), 0)
    total_heads = db.scalar("SELECT COUNT(*) FROM goat WHERE status='alive'", (), 0)
    span = max(1, (rules.parse_date(e) - rules.parse_date(s)).days + 1)

    feed = db.scalar(
        "SELECT COALESCE(SUM(t.amount),0) FROM txn t JOIN categories c ON c.id=t.category_id"
        " WHERE t.biz_date BETWEEN ? AND ? AND t.affect=1 AND c.group_name='饲料'", (s, e), 0) or 0

    per_head_rev = pl["total_income"] / (milk_heads or total_heads or 1) / span
    return {
        "start": s, "end": e,
        "milk": {
            "weight_kg": round(milk_kg, 1), "amount": round(milk_amt, 1),
            "days": days, "avg_daily": round(milk_kg / span, 1),
            "milking_heads": milk_heads,
            "per_head_day": round(milk_kg / span / milk_heads, 2) if milk_heads else 0,
            "avg_price": round(milk_amt / milk_kg, 2) if milk_kg else 0,
            "share": round(milk_amt / pl["total_income"] * 100, 1) if pl["total_income"] else 0,
        },
        "herd": {
            "total": total_heads, "milking": milk_heads,
            "per_head_revenue": round(per_head_rev, 1),
            "per_head_cost": round(pl["total_expense"] / (milk_heads or total_heads or 1) / span, 1),
        },
        "cost": {
            "total": pl["total_expense"], "feed": round(feed, 1),
            "feed_ratio": round(feed / pl["total_expense"] * 100, 1) if pl["total_expense"] else 0,
        },
        "profit": pl["profit"], "income": pl["total_income"], "expense": pl["total_expense"],
        "margin": pl["margin"],
    }


# ---------------- Excel 导出 ----------------
def _xlsx(sheets: list[tuple[str, list[list]]], title: str = "报表") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F6F4A")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    total_fill = PatternFill("solid", fgColor="EAF5EF")
    total_font = Font(bold=True)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for name, data in sheets:
        ws = wb.create_sheet(name[:31])
        if not data:
            continue
        for ri, rec in enumerate(data, start=1):
            for ci, val in enumerate(rec, start=1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.border = border
                if ri == 1:
                    cell.fill, cell.font = head_fill, head_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    if isinstance(val, (int, float)) and ri > 1:
                        cell.number_format = "#,##0.00"
                    if isinstance(rec[0], str) and ("合计" in rec[0] or "小计" in rec[0]
                                                   or "利润" in rec[0] or "总计" in rec[0]):
                        cell.fill, cell.font = total_fill, total_font
        widths = []
        for ci in range(1, len(data[0]) + 1):
            w = max((len(str(data[ri][ci - 1])) if ci - 1 < len(data[ri]) else 0)
                    for ri in range(len(data)))
            widths.append(min(max(w + 3, 10), 42))
        for ci, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _send(data: bytes, filename: str):
    from urllib.parse import quote
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":
                 f"attachment; filename=export.xlsx; filename*=UTF-8''{quote(filename)}"})


@router.get("/export/profit")
def export_profit(start: str = "", end: str = "", preset: str = "",
                  _: dict = Depends(require("rep.read"))):
    s, e = _range(start, end, preset)
    pl = _pl(s, e)
    data = [["科目类别", "明细科目", "金额(元)", "笔数"]]
    total_i = total_e = 0.0
    for g in pl["income"]:
        for i in g["items"]:
            total_i += i["amount"]
            data.append([g["group"], i["name"], i["amount"], i["count"]])
    data.append(["收入合计", "", round(total_i, 2), ""])
    data.append(["", "", "", ""])
    for g in pl["expense"]:
        for i in g["items"]:
            total_e += i["amount"]
            data.append([g["group"], i["name"], i["amount"], i["count"]])
    data.append(["支出合计", "", round(total_e, 2), ""])
    data.append(["", "", "", ""])
    data.append(["利润总额", "", pl["profit"], ""])
    data.append(["利润率%", "", pl["margin"], ""])
    cash = db.rows(
        "SELECT strftime('%Y-%m',biz_date) m,"
        " SUM(CASE WHEN direction='income' THEN amount ELSE 0 END) i,"
        " SUM(CASE WHEN direction='expense' THEN amount ELSE 0 END) x"
        " FROM txn WHERE biz_date BETWEEN ? AND ? AND affect=1 GROUP BY m", (s, e))
    monthly = [["月份", "收入", "支出", "利润"]]
    for r in cash:
        monthly.append([r["m"], round(r["i"] or 0, 2), round(r["x"] or 0, 2),
                        round((r["i"] or 0) - (r["x"] or 0), 2)])
    return _send(_xlsx([("利润表", data), ("逐月损益", monthly)]), f"利润表_{s}_{e}.xlsx")


@router.get("/export/txns")
def export_txns(start: str = "", end: str = "", preset: str = "",
                _: dict = Depends(require("fin.read"))):
    s, e = _range(start, end, preset)
    rs = db.rows(
        "SELECT t.biz_date, a.name acct, t.direction, c.name cat, c.group_name grp, t.amount,"
        " t.affect, t.counterparty, t.note, u.full_name usr"
        " FROM txn t JOIN accounts a ON a.id=t.account_id JOIN categories c ON c.id=t.category_id"
        " LEFT JOIN users u ON u.id=t.user_id WHERE t.biz_date BETWEEN ? AND ?"
        " ORDER BY t.biz_date, t.id", (s, e))
    data = [["日期", "账户", "方向", "科目", "科目分类", "金额", "计入损益", "对方", "摘要", "经手人"]]
    ti = te = 0.0
    for r in rs:
        data.append([r["biz_date"], r["acct"], "收" if r["direction"] == "income" else "支",
                     r["cat"], r["grp"], r["amount"], "是" if r["affect"] else "否",
                     r["counterparty"], r["note"], r["usr"]])
        if r["affect"]:
            if r["direction"] == "income":
                ti += r["amount"]
            else:
                te += r["amount"]
    data.append(["合计", "", "", "", "", round(ti - te, 2), "净损益", "", "", ""])
    return _send(_xlsx([("出入账明细", data)]), f"出入账明细_{s}_{e}.xlsx")


@router.get("/export/milk")
def export_milk(start: str = "", end: str = "", preset: str = "",
                _: dict = Depends(require("milk.read"))):
    s, e = _range(start, end, preset)
    rs = db.rows(
        "SELECT m.biz_date, m.group_code, m.batch, m.weight_kg, m.discard_kg, m.price, m.amount,"
        " m.grade, m.shift, m.note, u.full_name usr FROM milk m LEFT JOIN users u ON u.id=m.user_id"
        " WHERE m.biz_date BETWEEN ? AND ? ORDER BY m.biz_date, m.id", (s, e))
    data = [["日期", "分组", "批次", "产量(kg)", "弃奶(kg)", "单价", "金额", "等级", "班次", "备注", "录入人"]]
    tw = ta = td = 0.0
    for r in rs:
        data.append([r["biz_date"], r["group_code"], r["batch"], r["weight_kg"], r["discard_kg"],
                     r["price"], r["amount"], r["grade"], r["shift"], r["note"], r["usr"]])
        tw += r["weight_kg"]; ta += r["amount"]; td += r["discard_kg"] or 0
    data.append(["合计", "", "", round(tw, 2), round(td, 2), round(ta / tw, 2) if tw else 0,
                 round(ta, 2), "", "", "", ""])

    daily = db.rows(
        "SELECT biz_date, SUM(weight_kg) w, SUM(discard_kg) d, SUM(weight_kg*price) amt"
        " FROM milk WHERE biz_date BETWEEN ? AND ? GROUP BY biz_date ORDER BY biz_date", (s, e))
    dd = [["日期", "产量(kg)", "弃奶(kg)", "金额", "当日平均单产(kg/头·日)"]]
    for r in daily:
        h = db.scalar("SELECT COUNT(*) FROM goat WHERE status='alive' AND sex='f'"
                      " AND first_kidding IS NOT NULL AND birth_date<=? AND (status_date IS NULL OR status_date>?)",
                      (r["biz_date"], r["biz_date"]), 0)
        dd.append([r["biz_date"], round(r["w"] or 0, 2), round(r["d"] or 0, 2),
                   round(r["amt"] or 0, 2), round((r["w"] or 0) / h, 2) if h else ""])

    batch = db.rows(
        "SELECT batch, SUM(weight_kg) w, SUM(weight_kg*price) amt FROM milk"
        " WHERE biz_date BETWEEN ? AND ? GROUP BY batch", (s, e))
    bb = [["批次", "产量(kg)", "金额(元)", "占比%"]]
    for r in batch:
        bb.append([r["batch"], round(r["w"] or 0, 2), round(r["amt"] or 0, 2),
                   round((r["amt"] or 0) / ta * 100, 1) if ta else 0])
    return _send(_xlsx([("产奶明细", data), ("每日汇总", dd), ("批次汇总", bb)]), f"产奶记录_{s}_{e}.xlsx")


@router.get("/export/herd")
def export_herd(as_of: str = "", _: dict = Depends(require("goat.read"))):
    d = as_of or rules.today_str()
    rs = db.rows("SELECT * FROM goat WHERE status='alive' ORDER BY tag_no")
    data = [["耳标号", "性别", "分组", "日龄", "品种", "出生日期", "栋舍", "来源", "购入价",
             "首次产羔", "泌乳状态", "备注"]]
    from ..rules import group_of, GROUP_NAME, age_days
    for r in rs:
        g = group_of(r, d)
        data.append([r["tag_no"], "公" if r["sex"] == "m" else "母", GROUP_NAME.get(g, g),
                     age_days(r["birth_date"], d), r["breed"], r["birth_date"], r["barn"],
                     r["source"], r["buy_price"], r["first_kidding"] or "", r["milk_role"] or "", r["note"]])

    all_goats = db.rows("SELECT * FROM goat ORDER BY tag_no")
    full = [["耳标号", "状态", "性别", "品种", "出生日期", "离场日期", "离场原因", "离场收入",
             "栋舍", "母羊编号", "父羊编号", "备注"]]
    for r in all_goats:
        full.append([r["tag_no"], r["status"], "公" if r["sex"] == "m" else "母", r["breed"],
                     r["birth_date"], r["status_date"] or "", r["exit_reason"], r["exit_price"],
                     r["barn"], r["dam_tag"], r["sire_tag"], r["note"]])

    evs = db.rows(
        "SELECT e.event_date, g.tag_no, e.type, e.barn, e.weight, e.amount, e.milk_qty, e.note,"
        " u.full_name usr FROM goat_event e JOIN goat g ON g.id=e.goat_id"
        " LEFT JOIN users u ON u.id=e.user_id ORDER BY e.event_date DESC, e.id DESC LIMIT 20000")
    from ..rules import EVENT_TYPES
    ed = [["日期", "耳标号", "事件", "栋舍", "体重", "金额", "产奶", "备注", "经手人"]]
    for r in evs:
        ed.append([r["event_date"], r["tag_no"], EVENT_TYPES.get(r["type"], r["type"]), r["barn"],
                   r["weight"], r["amount"], r["milk_qty"], r["note"], r["usr"]])
    return _send(_xlsx([(f"在场存栏({d})", data), ("全部档案", full), ("事件流水", ed)]),
                 f"存栏档案_{d}.xlsx")


@router.get("/export/summary")
def export_summary(preset: str = "month", _: dict = Depends(require("rep.read"))):
    s, e = _range("", "", preset)
    bs = balance_sheet(as_of=e)
    cf = cashflow(start=s, end=e)
    pl = _pl(s, e)
    s1 = [["项目", "金额(元)", "说明"]]
    for a in bs["assets"]:
        s1.append([a["name"], a["amount"], a["note"]])
    s1.append(["资产总计", bs["totals"]["assets"], ""])
    for l in bs["liabilities"]:
        s1.append([l["name"], l["amount"], l["note"]])
    s1.append(["负债合计", bs["totals"]["liabilities"], ""])
    s1.append(["净资产", bs["totals"]["net_asset"], "资产 - 负债"])
    s1.append(["", "", ""])
    s1.append(["本年累计经营损益", bs["totals"]["year_profit"], f"{s[:4]} 年"])
    s1.append(["开业至今累计损益", bs["totals"]["cum_profit"], ""])
    s1.append(["在场存栏", bs["totals"]["heads"], "只"])

    s2 = [["月份", "流入", "流出", "净额"]]
    for r in cf["rows"]:
        s2.append([r["month"], r["inflow"], r["outflow"], r["net"]])
    s2.append(["合计", cf["total_inflow"], cf["total_outflow"], cf["net"]])

    s3 = [["科目分类", "明细科目", "金额", "笔数"]]
    for g in pl["income"]:
        for i in g["items"]:
            s3.append([g["group"], i["name"], i["amount"], i["count"]])
    s3.append(["收入合计", "", pl["total_income"], ""])
    for g in pl["expense"]:
        for i in g["items"]:
            s3.append([g["group"], i["name"], i["amount"], i["count"]])
    s3.append(["支出合计", "", pl["total_expense"], ""])
    s3.append(["利润", "", pl["profit"], ""])
    return _send(_xlsx([("资产负债简表", s1), ("现金流量表", s2), ("利润表", s3)]),
                 f"经营月报_{s}_{e}.xlsx")


@router.get("/export/summary-all")
def export_summary_all(preset: str = "year", _: dict = Depends(require("rep.read"))):
    return export_summary(preset or "year", _)
